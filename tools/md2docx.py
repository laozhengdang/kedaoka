#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把项目里的 Markdown 文档转成 Word（.docx）

只为这个项目的文档写的，支持的语法：
  # / ## / ### 标题，> 引用，- [ ] 复选框，| 表格 |，**粗体**，`代码`，[文字](链接)，
  --- 分隔线，普通段落
用法：
  python tools/md2docx.py 输入.md 输出.docx "文档标题"
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FONT_CN = "微软雅黑"
FONT_EN = "Segoe UI"
GRAY = RGBColor(0x66, 0x66, 0x66)
ACCENT = RGBColor(0x1F, 0x3A, 0x93)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_EN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT_EN)
    rfonts.set(qn("w:hAnsi"), FONT_EN)
    rfonts.set(qn("w:eastAsia"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_runs(par, text, size=10.5, base_bold=False, color=None):
    """处理 **粗体**、`代码`、[文字](链接)"""
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    for seg in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = par.add_run(seg[2:-2])
            set_run_font(r, size, True, color)
        elif seg.startswith("`") and seg.endswith("`"):
            r = par.add_run(seg[1:-1])
            set_run_font(r, size - 0.5, base_bold, RGBColor(0xA0, 0x30, 0x30))
        else:
            r = par.add_run(seg)
            set_run_font(r, size, base_bold, color)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_hr(doc):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "CCCCCC")
    borders.append(bottom)
    ppr.append(borders)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)


def flush_table(doc, rows):
    if not rows:
        return
    header = rows[0]
    body = rows[1:]
    cols = len(header)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.space_before = Pt(1)
            val = row[j] if j < len(row) else ""
            if i == 0:
                add_runs(par, val, 9.5, True)
                shade(cell, "EEF2FF")
            else:
                # 空单元格留空，方便手写
                add_runs(par, val, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path: Path, out_path: Path, title_hint: str | None = None):
    lines = md_path.read_text(encoding="utf-8").split("\n")
    doc = Document()

    # 页面：A4、窄一点边距
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = FONT_EN
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    table_rows = []
    i = 0
    first_heading_done = False
    in_code = False
    code_buf = []

    def flush_code():
        nonlocal code_buf
        if not code_buf:
            return
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.4)
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(6)
        ppr = par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F4F4F4")
        ppr.append(shd)
        for idx, cl in enumerate(code_buf):
            r = par.add_run(cl)
            r.font.name = "Consolas"
            rpr = r._element.get_or_add_rPr()
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), "Consolas")
            rfonts.set(qn("w:hAnsi"), "Consolas")
            rfonts.set(qn("w:eastAsia"), FONT_CN)
            rpr.append(rfonts)
            r.font.size = Pt(9)
            if idx < len(code_buf) - 1:
                r.add_break()
        code_buf = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # 围栏代码块
        if line.strip().startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        # 表格聚合
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue  # 分隔行
            table_rows.append(cells)
            i += 1
            continue
        elif table_rows:
            flush_table(doc, table_rows)
            table_rows = []

        if not line.strip():
            i += 1
            continue

        if line.startswith("---"):
            add_hr(doc)
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", m.group(2))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if level == 1 else (10 if level == 2 else 8))
            p.paragraph_format.space_after = Pt(4)
            if level == 1 and not first_heading_done:
                first_heading_done = True
                add_runs(p, text, 20, True, ACCENT)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif level == 1:
                add_runs(p, text, 16, True, ACCENT)
            elif level == 2:
                add_runs(p, text, 13, True, RGBColor(0x22, 0x22, 0x22))
            else:
                add_runs(p, text, 11.5, True, RGBColor(0x33, 0x33, 0x33))
            i += 1
            continue

        if line.startswith(">"):
            text = line.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, text, 9.5, False, GRAY)
            i += 1
            continue

        m = re.match(r"^-\s+\[( |x|X)\]\s*(.*)$", line)
        if m:
            checked = m.group(1).lower() == "x"
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, ("☑ " if checked else "☐ ") + m.group(2), 10.5)
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, "· " + m.group(1), 10.5)
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, f"{m.group(1)}. {m.group(2)}", 10.5)
            i += 1
            continue

        # 「答：」「答:」这类填写行：后面补一条可书写的下划线
        m_ans = re.match(r"^(答[：:])\s*(.*)$", line.strip())
        if m_ans:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.35
            r0 = p.add_run(m_ans.group(1))
            set_run_font(r0, 10.5, True)
            if m_ans.group(2):
                add_runs(p, " " + m_ans.group(2), 10.5)
            else:
                blank = p.add_run(" " + "\u3000" * 22)
                set_run_font(blank, 10.5)
            # 段落底边框当书写线（比下划线可靠）
            ppr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "dotted")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "999999")
            borders.append(bottom)
            ppr.append(borders)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        add_runs(p, line.strip(), 10.5)
        i += 1

    if table_rows:
        flush_table(doc, table_rows)

    # 页脚
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"{title_hint or md_path.stem} · 老郑荡的妙妙小工具")
    set_run_font(r, 8, False, GRAY)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("用法: python tools/md2docx.py 输入.md 输出.docx [标题]")
        sys.exit(1)
    out = convert(Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3] if len(sys.argv) > 3 else None)
    print(f"已生成: {out}  ({out.stat().st_size/1024:.1f} KB)")
